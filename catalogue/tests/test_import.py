import pytest
from django.core.management import CommandError, call_command
from openpyxl import Workbook, load_workbook

from catalogue.models import Produit
from catalogue.tests.factories import ProduitFactory
from comptes.tests.factories import UtilisateurFactory
from entreprises.tests.factories import EntrepriseFactory

pytestmark = pytest.mark.django_db

ENTETES = ["Code", "Nom", "Unité", "Prix", "Notes préparation", "Actif"]


def classeur(chemin, lignes, entetes=ENTETES):
    wb = Workbook()
    wb.active.append(entetes)
    for ligne in lignes:
        wb.active.append(ligne)
    wb.save(chemin)
    return chemin


def importer(chemin, entreprise, *options):
    from io import StringIO

    sortie = StringIO()
    call_command(
        "importer_produits", str(chemin), "--entreprise", entreprise.slug, *options, stdout=sortie
    )
    return sortie.getvalue()


@pytest.fixture
def entreprise():
    return EntrepriseFactory(slug="mie-dore")


def produits(entreprise):
    return {p.code: p for p in Produit.toutes_entreprises.filter(entreprise=entreprise)}


def test_import_excel(tmp_path, entreprise):
    fichier = classeur(
        tmp_path / "produits.xlsx",
        [
            ["BAG", "Pain baguette", "pièce", 250, "", "oui"],
            ["CRO", "Croissant", "Piece", 300.0, "Beurre AOP", ""],
            ["FAR", "Farine", "kg", "1 200", "", "non"],
        ],
    )
    sortie = importer(fichier, entreprise)
    assert "3 créé(s), 0 mis à jour, 0 ignoré(s), 0 rejeté(s)." in sortie
    tous = produits(entreprise)
    assert tous["CRO"].prix_reference == 300 and tous["CRO"].notes_preparation == "Beurre AOP"
    assert tous["FAR"].unite == "kg" and tous["FAR"].prix_reference == 1200
    assert tous["FAR"].actif is False


def test_import_csv_point_virgule(tmp_path, entreprise):
    fichier = tmp_path / "produits.csv"
    fichier.write_text(
        "code;nom;unite;prix\nBAG;Pain baguette;pièce;250\nMIE;Pain de mie;carton;700\n",
        encoding="utf-8-sig",
    )
    assert "2 créé(s)" in importer(fichier, entreprise)


def test_lignes_rejetees_sans_bloquer_les_autres(tmp_path, entreprise):
    fichier = classeur(
        tmp_path / "produits.xlsx",
        [
            ["BAG", "Pain baguette", "pièce", 250, "", ""],
            ["", "Sans code", "pièce", 100, "", ""],
            ["SAC", "Sachet", "sachet", 50, "", ""],
            ["CRO", "Croissant", "pièce", "12,5", "", ""],
            ["NEG", "Négatif", "pièce", -5, "", ""],
        ],
    )
    sortie = importer(fichier, entreprise)
    assert "1 créé(s), 0 mis à jour, 0 ignoré(s), 4 rejeté(s)." in sortie
    assert "Ligne 3 rejetée : code manquant." in sortie
    assert "Ligne 4 rejetée : Unité « sachet » inconnu(e)" in sortie
    assert "Ligne 5 rejetée : Prix « 12,5 » doit être un nombre entier." in sortie
    assert "Ligne 6 rejetée : prix_reference" in sortie
    assert list(produits(entreprise)) == ["BAG"]


def test_relance_sans_doublon_et_mise_a_jour(tmp_path, entreprise):
    lignes = [["BAG", "Pain baguette", "pièce", 250, "", ""]]
    importer(classeur(tmp_path / "a.xlsx", lignes), entreprise)
    assert "0 créé(s), 0 mis à jour, 1 ignoré(s)" in importer(
        classeur(tmp_path / "b.xlsx", lignes), entreprise
    )

    lignes = [["BAG", "Baguette", "pièce", 275, "", "non"]]
    assert "1 mis à jour" in importer(classeur(tmp_path / "c.xlsx", lignes), entreprise)
    bag = produits(entreprise)["BAG"]
    assert (bag.nom, bag.prix_reference, bag.actif) == ("Baguette", 275, False)


def test_simulation_n_enregistre_rien(tmp_path, entreprise):
    fichier = classeur(tmp_path / "produits.xlsx", [["BAG", "Pain baguette", "pièce", 250, "", ""]])
    sortie = importer(fichier, entreprise, "--simulation")
    assert "Simulation (rien n'est enregistré) : 1 créé(s)" in sortie
    assert not Produit.toutes_entreprises.exists()


def test_import_limite_a_l_entreprise_choisie(tmp_path, entreprise):
    autre = EntrepriseFactory()
    ProduitFactory(entreprise=autre, code="BAG", nom="Baguette d'ailleurs", prix_reference=999)
    importer(
        classeur(tmp_path / "p.xlsx", [["BAG", "Pain baguette", "pièce", 250, "", ""]]), entreprise
    )
    assert produits(entreprise)["BAG"].prix_reference == 250
    assert produits(autre)["BAG"].prix_reference == 999


def test_auteur_de_l_import(tmp_path, entreprise):
    aminata = UtilisateurFactory(entreprise=entreprise, username="aminata.k")
    fichier = classeur(tmp_path / "p.xlsx", [["BAG", "Pain baguette", "pièce", 250, "", ""]])
    importer(fichier, entreprise, "--utilisateur", "aminata.k")
    assert produits(entreprise)["BAG"].cree_par == aminata


@pytest.mark.parametrize(
    ("options", "message"),
    [
        (["--entreprise", "inconnue"], "Entreprise « inconnue » introuvable"),
        ([], "Préciser l'entreprise"),
        (["--entreprise", "mie-dore", "--utilisateur", "personne"], "Utilisateur « personne »"),
    ],
)
def test_erreurs_de_lancement(tmp_path, entreprise, options, message):
    fichier = classeur(tmp_path / "p.xlsx", [])
    with pytest.raises(CommandError, match=message):
        call_command("importer_produits", str(fichier), *options)


def test_colonne_obligatoire_manquante(tmp_path, entreprise):
    fichier = classeur(tmp_path / "p.xlsx", [["BAG", "Pain baguette"]], entetes=["code", "nom"])
    with pytest.raises(CommandError, match="Colonnes manquantes : unite, prix"):
        importer(fichier, entreprise)


def test_modele_de_fichier(tmp_path):
    chemin = tmp_path / "modele.xlsx"
    call_command("importer_produits", str(chemin), "--modele")
    lignes = list(load_workbook(chemin).active.iter_rows(values_only=True))
    assert lignes[0] == ("code", "nom", "unite", "prix", "notes_preparation", "actif")
    assert lignes[1][:2] == ("BAG", "Pain baguette")
