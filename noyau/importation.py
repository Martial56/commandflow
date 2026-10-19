"""Import de fichiers Excel (.xlsx) ou CSV, ligne par ligne, à travers les services métier.

Chaque ligne est traitée dans son propre point de sauvegarde : une ligne rejetée
n'annule pas les autres. En simulation, tout est annulé à la fin.
"""

import csv
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook


class LigneRejetee(Exception):
    """Ligne invalide : le message est rapporté tel quel."""


def normaliser(valeur):
    """« Fin de mois » → « fin de mois » ; « Pièce » → « piece » (comparaisons souples)."""
    texte = unicodedata.normalize("NFKD", str(valeur or "")).encode("ascii", "ignore").decode()
    return " ".join(texte.casefold().replace("_", " ").split())


def _entete(valeur):
    return normaliser(valeur).replace(" ", "_")


def lire_tableau(chemin):
    """Lignes du fichier sous forme de dictionnaires {en-tête normalisé: valeur}.

    Renvoie des couples (numéro de ligne dans le fichier, données). Les lignes vides
    sont ignorées. Le CSV peut utiliser « ; » (Excel français) ou « , ».
    """
    chemin = Path(chemin)
    if chemin.suffix.lower() == ".xlsx":
        feuille = load_workbook(chemin, read_only=True, data_only=True).active
        lignes = [list(ligne) for ligne in feuille.iter_rows(values_only=True)]
    elif chemin.suffix.lower() == ".csv":
        with chemin.open(encoding="utf-8-sig", newline="") as fichier:
            echantillon = fichier.read(4096)
            fichier.seek(0)
            dialecte = csv.Sniffer().sniff(echantillon, delimiters=";,")
            lignes = list(csv.reader(fichier, dialecte))
    else:
        raise ValueError("Format non pris en charge : utiliser un fichier .xlsx ou .csv.")

    if not lignes:
        return []
    entetes = [_entete(valeur) for valeur in lignes[0]]
    resultat = []
    for numero, valeurs in enumerate(lignes[1:], start=2):
        donnees = {
            entete: ("" if valeur is None else str(valeur).strip())
            for entete, valeur in zip(entetes, valeurs, strict=False)
            if entete
        }
        if any(donnees.values()):
            resultat.append((numero, donnees))
    return resultat


def ecrire_modele(chemin, colonnes, exemple):
    """Fichier Excel modèle : une ligne d'en-têtes et une ligne d'exemple."""
    classeur = Workbook()
    feuille = classeur.active
    feuille.append(colonnes)
    feuille.append([exemple.get(colonne, "") for colonne in colonnes])
    for cellule in feuille[1]:
        feuille.column_dimensions[cellule.column_letter].width = max(14, len(cellule.value) + 4)
    classeur.save(chemin)


def choix(valeur, options, nom_colonne):
    """Valeur d'un `TextChoices` à partir de sa valeur ou de son libellé."""
    cle = normaliser(valeur)
    for option in options:
        if cle in (normaliser(option.value), normaliser(option.label)):
            return option.value
    attendues = ", ".join(str(option.label) for option in options)
    raise LigneRejetee(f"{nom_colonne} « {valeur} » inconnu(e) ; valeurs possibles : {attendues}.")


def nombre(valeur, nom_colonne, *, entier=False, obligatoire=False):
    """« 1 250 », « 1250,0 » ou 1250.0 (Excel) → Decimal. Vide → None."""
    texte = str(valeur or "").replace(" ", "").replace("\xa0", "").replace(" ", "")
    if not texte:
        if obligatoire:
            raise LigneRejetee(f"{nom_colonne} manquant(e).")
        return None
    try:
        resultat = Decimal(texte.replace(",", "."))
    except InvalidOperation:
        raise LigneRejetee(f"{nom_colonne} « {valeur} » n'est pas un nombre.") from None
    if entier:
        if resultat != resultat.to_integral_value():
            raise LigneRejetee(f"{nom_colonne} « {valeur} » doit être un nombre entier.")
        resultat = resultat.quantize(Decimal(1))
    return resultat


def oui_non(valeur, *, defaut=True):
    texte = normaliser(valeur)
    if not texte:
        return defaut
    if texte in {"oui", "o", "1", "vrai", "x"}:
        return True
    if texte in {"non", "n", "0", "faux"}:
        return False
    raise LigneRejetee(f"« {valeur} » : indiquer oui ou non.")


def message_erreur(erreur):
    if isinstance(erreur, ValidationError):
        if hasattr(erreur, "message_dict"):
            return " ; ".join(
                f"{champ} : {' '.join(messages)}" if champ != "__all__" else " ".join(messages)
                for champ, messages in erreur.message_dict.items()
            )
        return " ".join(erreur.messages)
    return str(erreur)


@dataclass
class Rapport:
    crees: int = 0
    mis_a_jour: int = 0
    ignores: int = 0
    rejets: list = field(default_factory=list)

    def resume(self):
        return (
            f"{self.crees} créé(s), {self.mis_a_jour} mis à jour, {self.ignores} ignoré(s), "
            f"{len(self.rejets)} rejeté(s)."
        )


class Simulation(Exception):
    """Annule toute la transaction d'un import en mode simulation."""


def importer(lignes, traiter_ligne, *, simulation=False):
    """Applique `traiter_ligne(donnees)` à chaque ligne.

    `traiter_ligne` renvoie « cree », « mis_a_jour » ou « ignore », ou lève
    `LigneRejetee`, `ValidationError` ou `ValueError` pour rejeter la ligne.
    """
    rapport = Rapport()
    try:
        with transaction.atomic():
            for numero, donnees in lignes:
                try:
                    with transaction.atomic():
                        resultat = traiter_ligne(donnees)
                except (LigneRejetee, ValidationError, ValueError) as erreur:
                    rapport.rejets.append((numero, message_erreur(erreur)))
                    continue
                if resultat == "cree":
                    rapport.crees += 1
                elif resultat == "mis_a_jour":
                    rapport.mis_a_jour += 1
                else:
                    rapport.ignores += 1
            if simulation:
                raise Simulation
    except Simulation:
        pass
    return rapport
