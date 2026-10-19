from io import StringIO

import pytest
from django.core.management import call_command
from openpyxl import Workbook, load_workbook

from clients.models import Client
from clients.tests.factories import ClientFactory
from entreprises.tests.factories import EntrepriseFactory

pytestmark = pytest.mark.django_db

ENTETES = [
    "code", "nom", "catégorie", "téléphone", "conditions_paiement", "delai_paiement_jours",
    "adresse_libelle", "adresse", "commune", "contact_nom", "contact_fonction",
]  # fmt: skip
BAOBAB = ["", "Restaurant Le Baobab", "Restaurant", "07 00 00 00 01", "Fin de mois", "",
          "Cuisine", "Face pharmacie", "Cocody", "M. Koffi", "gérant"]  # fmt: skip
PALMIERS = ["", "Supérette Les Palmiers", "supermarché", "05 00 00 00 02", "", "30",
            "", "Rue 12", "Marcory", "", ""]  # fmt: skip


def classeur(chemin, lignes):
    wb = Workbook()
    wb.active.append(ENTETES)
    for ligne in lignes:
        wb.active.append(ligne)
    wb.save(chemin)
    return chemin


def importer(chemin, entreprise, *options):
    sortie = StringIO()
    call_command(
        "importer_clients", str(chemin), "--entreprise", entreprise.slug, *options, stdout=sortie
    )
    return sortie.getvalue()


@pytest.fixture
def entreprise():
    return EntrepriseFactory(slug="mie-dore")


def clients(entreprise):
    return {c.nom: c for c in Client.toutes_entreprises.filter(entreprise=entreprise)}


def test_creation_avec_adresse_et_contact(tmp_path, entreprise):
    sortie = importer(classeur(tmp_path / "c.xlsx", [BAOBAB, PALMIERS]), entreprise)
    assert "2 créé(s), 0 mis à jour, 0 ignoré(s), 0 rejeté(s)." in sortie
    tous = clients(entreprise)
    baobab, palmiers = tous["Restaurant Le Baobab"], tous["Supérette Les Palmiers"]
    assert {baobab.code, palmiers.code} == {"CLI-000001", "CLI-000002"}
    assert baobab.conditions_paiement == "fin_de_mois"
    adresse = baobab.adresses.model.toutes_entreprises.get(client=baobab)
    assert (adresse.libelle, adresse.commune, adresse.par_defaut) == ("Cuisine", "Cocody", True)
    assert baobab.contacts.model.toutes_entreprises.get(client=baobab).fonction == "gérant"
    # Délai renseigné sans conditions : paiement à X jours ; libellé d'adresse par défaut.
    assert (palmiers.conditions_paiement, palmiers.delai_paiement_jours) == ("x_jours", 30)
    assert palmiers.adresses.model.toutes_entreprises.get(client=palmiers).libelle == (
        "Adresse principale"
    )


def test_relance_sans_doublon(tmp_path, entreprise):
    importer(classeur(tmp_path / "a.xlsx", [BAOBAB]), entreprise)
    sortie = importer(classeur(tmp_path / "b.xlsx", [BAOBAB]), entreprise)
    assert "0 créé(s), 0 mis à jour, 1 ignoré(s)" in sortie
    assert len(clients(entreprise)) == 1


def test_mise_a_jour_par_code_sans_effacer(tmp_path, entreprise):
    importer(classeur(tmp_path / "a.xlsx", [BAOBAB]), entreprise)
    ligne = ["CLI-000001", "", "Restaurant", "", "", "15", "", "", "", "", ""]
    assert "1 mis à jour" in importer(classeur(tmp_path / "b.xlsx", [ligne]), entreprise)
    baobab = clients(entreprise)["Restaurant Le Baobab"]
    assert (baobab.conditions_paiement, baobab.delai_paiement_jours) == ("x_jours", 15)
    assert baobab.telephone == "07 00 00 00 01"


def test_lignes_rejetees(tmp_path, entreprise):
    lignes = [
        BAOBAB,
        ["", "Sans catégorie", "", "07 11 11 11 11", "", "", "", "", "", "", ""],
        ["", "Catégorie inconnue", "Hôtel", "07 22 22 22 22", "", "", "", "", "", "", ""],
        ["", "Sans commune", "Restaurant", "07 33 33 33 33", "", "", "", "Rue 1", "", "", ""],
        ["CLI-000999", "Code inconnu", "Restaurant", "07 44 44 44 44", "", "", "", "", "", "", ""],
        ["", "Sans téléphone", "Particulier", "", "", "", "", "", "", "", ""],
    ]
    sortie = importer(classeur(tmp_path / "c.xlsx", lignes), entreprise)
    assert "1 créé(s), 0 mis à jour, 0 ignoré(s), 5 rejeté(s)." in sortie
    assert "Ligne 3 rejetée : Catégorie «  » inconnu(e)" in sortie
    assert "Ligne 4 rejetée : Catégorie « Hôtel » inconnu(e)" in sortie
    assert "Ligne 5 rejetée : commune" in sortie
    assert "Ligne 6 rejetée : code client « CLI-000999 » inconnu" in sortie
    assert "Ligne 7 rejetée : telephone" in sortie
    # Les numéros des lignes rejetées ne sont pas consommés.
    assert clients(entreprise)["Restaurant Le Baobab"].code == "CLI-000001"
    importer(classeur(tmp_path / "d.xlsx", [PALMIERS]), entreprise)
    assert clients(entreprise)["Supérette Les Palmiers"].code == "CLI-000002"


def test_doublon_d_une_autre_entreprise_ignore(tmp_path, entreprise):
    ClientFactory(
        entreprise=EntrepriseFactory(), nom="Restaurant Le Baobab", telephone="07 00 00 00 01"
    )
    assert "1 créé(s)" in importer(classeur(tmp_path / "c.xlsx", [BAOBAB]), entreprise)


def test_simulation(tmp_path, entreprise):
    sortie = importer(classeur(tmp_path / "c.xlsx", [BAOBAB, PALMIERS]), entreprise, "--simulation")
    assert "Simulation (rien n'est enregistré) : 2 créé(s)" in sortie
    assert not Client.toutes_entreprises.exists()


def test_modele_de_fichier(tmp_path):
    chemin = tmp_path / "modele.xlsx"
    call_command("importer_clients", str(chemin), "--modele")
    entetes, exemple = list(load_workbook(chemin).active.iter_rows(values_only=True))
    assert entetes[:4] == ("code", "nom", "categorie", "telephone")
    assert exemple[1] == "Restaurant Le Baobab"
